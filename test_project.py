"""Run with python -m unittest discover -v from utilities."""
import sys
import unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from utilities.document import corrected_bytes
from utilities.analysis_service import analyze
from utilities.font_and_styles_handling import check_font_size, check_caption_styles
from utilities.pagenumber import is_page_number_field
from utilities.spacing_handling import check_line_spacing


def sample_bytes():
    doc = Document()
    doc.add_paragraph('Research <methods> & results').runs[0].font.size = Pt(20)
    doc.add_paragraph('Figure 1: Findings')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Preserved table'
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class DocumentTests(unittest.TestCase):
    def test_correction_preserves_content_and_matches_validator(self):
        data = corrected_bytes(sample_bytes())
        doc = Document(BytesIO(data))
        self.assertEqual(doc.paragraphs[0].text, 'Research <methods> & results')
        self.assertEqual(doc.tables[0].cell(0, 0).text, 'Preserved table')
        self.assertFalse(any('Error:' in line for line in check_font_size(doc)))
        self.assertFalse(any('Error:' in line for line in check_caption_styles(doc)))
        self.assertFalse(any('Error:' in line for line in check_line_spacing(doc)))
        again = Document(BytesIO(corrected_bytes(data)))
        self.assertEqual(doc._element.xml, again._element.xml)

    def test_plain_text_and_numpages_are_not_page_fields(self):
        para = Document().add_paragraph('PAGE')
        self.assertFalse(is_page_number_field(para))
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), ' NUMPAGES ')
        para._p.append(field)
        self.assertFalse(is_page_number_field(para))
        field.set(qn('w:instr'), ' PAGE ')
        self.assertTrue(is_page_number_field(para))

    def test_inherited_font_size_is_checked(self):
        doc = Document()
        doc.styles['Normal'].font.size = Pt(20)
        doc.add_paragraph('Inherited size')
        self.assertTrue(any('Error:' in line for line in check_font_size(doc)))

    def test_analysis_produces_all_downloads_without_ai_call(self):
        with patch('utilities.analysis_service.check_ai_plagiarism') as detector:
            result = analyze(sample_bytes())
            detector.assert_not_called()
        self.assertTrue(result['pdf'].startswith(b'%PDF'))
        Document(BytesIO(result['highlighted']))
        Document(BytesIO(result['corrected']))

    def test_invalid_upload_is_rejected(self):
        with self.assertRaises(ValueError):
            analyze(b'not a word document')


if __name__ == '__main__':
    unittest.main()
