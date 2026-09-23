"""Version 2: profiles, approval boundaries, document evidence and local/AI chat."""
import json
import os
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch, Mock
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document
from docx.shared import Pt, Inches
from fastapi.testclient import TestClient
from utilities.api import create_app
from utilities.profiles import FormattingProfile
from utilities.document import corrected_bytes
from utilities.review_engine import inspect_document, readiness
from utilities.corrections import preview
from utilities.document_assistant import ai_answer, local_answer


def sample():
    doc = Document()
    doc.add_heading('Introduction', 1)
    doc.add_paragraph('Our solar research examines renewable energy in rural areas.').runs[0].font.size = Pt(20)
    doc.add_heading('Methodology', 1)
    doc.add_paragraph('We surveyed a sample of households. This document is an academic draft.')
    doc.add_table(rows=1, cols=1).cell(0,0).text = 'Table evidence must remain unchanged.'
    buf = BytesIO(); doc.save(buf)
    return buf.getvalue()


class ProfileAndCorrectionTests(unittest.TestCase):
    def test_selected_rules_agree_with_correction(self):
        profile = FormattingProfile(body_font='Times New Roman',body_size=13,heading_size=16,line_spacing=2,margin_left=1.5)
        original = Document(BytesIO(sample()))
        corrected = Document(BytesIO(corrected_bytes(sample(),profile)))
        findings = inspect_document(corrected,profile)['findings']
        self.assertFalse([f for f in findings if f['correction_group']])
        self.assertEqual([p.text for p in original.paragraphs],[p.text for p in corrected.paragraphs])
        self.assertEqual(original.tables[0].cell(0,0).text,corrected.tables[0].cell(0,0).text)
        self.assertAlmostEqual(corrected.sections[0].left_margin.inches,1.5)

    def test_selective_changes_preserve_rejected_categories(self):
        original = Document(BytesIO(sample()))
        plan, output = preview(sample(),FormattingProfile(),['margins'])
        result = Document(BytesIO(output))
        self.assertEqual(original.paragraphs[1]._p.xml,result.paragraphs[1]._p.xml)
        self.assertTrue(all(c['group']=='margins' for c in plan['changes']))
        changed = preview(sample(),FormattingProfile(),['sizes'])[0]
        self.assertNotEqual(plan['preview_id'],changed['preview_id'])

    def test_readiness_has_evidence_and_no_automatic_pass(self):
        review = inspect_document(Document(BytesIO(sample())))
        report = readiness(review)
        tasks = {t['title']:t for t in report['tasks']}
        self.assertEqual(tasks['Introduction section']['location'],'P1')
        self.assertEqual(tasks['Abstract section']['status'],'needs_review')
        self.assertEqual(tasks['Academic and citation review']['status'],'manual_required')

    def test_local_assistant_cites_paragraph_and_cannot_apply_changes(self):
        review = inspect_document(Document(BytesIO(sample())))
        review['readiness'] = readiness(review)
        answer = local_answer(review,'Explain formatting errors in paragraph P2')
        self.assertIn('[P2]',answer['answer'])
        self.assertEqual(answer['citations'][0]['id'],'P2')
        self.assertEqual(local_answer(review,'Fix my headings')['action'],'corrections')

    @patch.dict(os.environ,{'OPENAI_API_KEY':'test-key','ADF_AI_MODEL':'test-model'})
    @patch('utilities.document_assistant.requests.post')
    def test_ai_uses_bounded_evidence_and_no_tools(self, post):
        review = inspect_document(Document(BytesIO(sample())))
        review['readiness'] = readiness(review)
        post.return_value = Mock(raise_for_status=lambda:None,json=lambda:{'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':'This passage discusses solar research [P2].'}]}]})
        answer = ai_answer(review,'What does paragraph P2 discuss?',[])
        self.assertEqual(answer['mode'],'ai')
        payload = post.call_args.kwargs['json']
        self.assertFalse(payload['store'])
        self.assertNotIn('tools',payload)
        self.assertEqual(answer['citations'][0]['id'],'P2')
        post.return_value.json = lambda:{'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':'Invented [P999].'}]}]}
        with self.assertRaises(ValueError):
            ai_answer(review,'Explain P2',[])


class V2APITests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.app=create_app(self.temp.name)
        self.client=TestClient(self.app,headers={'X-Requested-With':'ADF-Web'})
        self.client.__enter__()
        self.client.post('/api/auth/signup',json={'username':'researcher','password':'strong-password'})

    def tearDown(self):
        self.client.__exit__(None,None,None)
        self.temp.cleanup()

    def upload(self,profile_id='default'):
        response=self.client.post('/api/reviews?profile_id='+profile_id,content=sample(),headers={'X-Filename':'sample.docx'})
        self.assertEqual(response.status_code,201,response.text)
        return response.json()

    def test_profile_saved_and_snapshot_survives_profile_deletion(self):
        result=self.client.post('/api/profiles',json={'name':'My University handbook','body_font':'Times New Roman','line_spacing':2})
        self.assertEqual(result.status_code,201,result.text)
        profile=result.json()
        review=self.upload(profile['id'])
        self.assertEqual(review['profile']['line_spacing'],2)
        self.client.delete('/api/profiles/'+profile['id'])
        self.assertEqual(self.client.get('/api/reviews/'+review['id']).json()['profile']['body_font'],'Times New Roman')
        self.assertEqual(self.client.post('/api/profiles',json={'margin_left':99}).status_code,422)

    def test_approval_is_required_and_bound_to_selection(self):
        review=self.upload();base='/api/reviews/'+review['id']
        self.assertEqual(self.client.get(base+'/download/corrected').status_code,409)
        plan=self.client.post(base+'/corrections/preview',json={'groups':['margins']}).json()
        changed=self.client.post(base+'/corrections/approve',json={'groups':['fonts'],'preview_id':plan['preview_id'],'approved':True})
        self.assertEqual(changed.status_code,409)
        denied=self.client.post(base+'/corrections/approve',json={'groups':['margins'],'preview_id':plan['preview_id'],'approved':False})
        self.assertEqual(denied.status_code,422)
        approved=self.client.post(base+'/corrections/approve',json={'groups':['margins'],'preview_id':plan['preview_id'],'approved':True})
        self.assertEqual(approved.status_code,200,approved.text)
        self.assertIn('remaining_findings',approved.json())
        output=self.client.get(base+'/download/corrected')
        self.assertEqual(output.status_code,200)
        self.assertEqual(Document(BytesIO(output.content)).paragraphs[1].runs[0].font.size.pt,20)

    @patch('utilities.document_assistant.requests.post')
    def test_local_chat_never_calls_external_service_and_acknowledgement_is_not_pass(self, post):
        review=self.upload();base='/api/reviews/'+review['id']
        response=self.client.post(base+'/assistant',json={'question':'Explain issues in P2'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['mode'],'local')
        post.assert_not_called()
        self.assertEqual(self.client.post(base+'/assistant',json={'question':'Explain P2','use_ai':True,'consent':False}).status_code,422)
        task=review['readiness']['tasks'][0]['id']
        checked=self.client.post(base+'/readiness/acknowledge',json={'task_id':task,'checked':True})
        self.assertIn(task,checked.json()['acknowledged_tasks'])
        self.assertEqual(self.client.get(base).json()['total'],review['total'])

    def test_new_endpoints_are_owner_scoped(self):
        profile=self.client.post('/api/profiles',json={'name':'Private rules'}).json()
        review=self.upload(profile['id']);base='/api/reviews/'+review['id']
        self.client.post('/api/auth/logout')
        self.client.post('/api/auth/signup',json={'username':'another-user','password':'strong-password'})
        self.assertEqual(self.client.get('/api/profiles').json()[0]['id'],'default')
        self.assertEqual(len(self.client.get('/api/profiles').json()),1)
        self.assertEqual(self.client.get(base+'/paragraphs/P2').status_code,404)
        self.assertEqual(self.client.post(base+'/assistant',json={'question':'Explain P2'}).status_code,404)
        self.assertEqual(self.client.post(base+'/corrections/preview',json={'groups':['fonts']}).status_code,404)
        self.assertEqual(self.client.delete('/api/profiles/'+profile['id']).status_code,404)


if __name__=='__main__':
    unittest.main()
